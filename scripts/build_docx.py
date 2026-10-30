"""Build editable Word (.docx) files of the Fisko SRS and SDS from the same parsed blocks as the HTML.

Usage
    python scripts/build_docx.py

Outputs (built, not committed): docs/spec/dist/Fisko-{SRS,SDS}-{th,en}.docx
"""
from __future__ import annotations

import functools
import io
import json
import pathlib
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import build_spec as bs  # noqa: E402

ROOT = bs.ROOT
DIAG_PNG = bs.SPEC / 'diagrams' / 'png'

THAI_FONT = 'TH Sarabun New'
LATIN_FONT = 'Calibri'
MONO_FONT = 'Consolas'
GREEN = RGBColor(0x16, 0x65, 0x34)
TEXT_W = Cm(16.0)       # A4 with 2.5 cm left and 2.5 cm right margins
MAX_FIG_H = Cm(21.0)

# body size, table size and heading sizes per language: Thai fonts read smaller at the same point size
SIZES = {'th': {'body': 16, 'table': 14, 'code': 9, 'h': [22, 19, 17, 16]},
         'en': {'body': 10.5, 'table': 9, 'code': 8.5, 'h': [16, 13.5, 12, 11]}}

CALLOUT_FILL = {'NOTE': 'F4F5F7', 'WARN': 'FBF6EA', 'OPEN': 'ECF0FA', 'LAW': 'E8F5EC'}



# ----------------------------------------------------------------------------------------------------------
# low-level helpers
# ----------------------------------------------------------------------------------------------------------

def cs_size(lang: str, size: float) -> float:
    """Thai is a complex script in Word with its own size: keep Thai readable inside the English file too."""
    return size if lang == 'th' else size * 1.5


def style_font(style, latin: str, size: float, cs: float) -> None:
    """Latin and complex-script (Thai) font and size on a style, without theme fonts overriding them."""
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for attr in list(fonts.attrib):
        if attr.lower().endswith('theme'):
            del fonts.attrib[attr]
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
        fonts.set(qn(attr), latin)
    fonts.set(qn('w:cs'), THAI_FONT)
    szcs = rpr.find(qn('w:szCs'))
    if szcs is None:
        szcs = OxmlElement('w:szCs')
        rpr.append(szcs)
    szcs.set(qn('w:val'), str(int(round(cs * 2))))


def run_size(run, lang: str, size: float) -> None:
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    szcs = OxmlElement('w:szCs')
    szcs.set(qn('w:val'), str(int(round(cs_size(lang, size) * 2))))
    rpr.append(szcs)


def shade(cell, fill: str) -> None:
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcpr.append(shd)


def repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    trpr.append(el)


# style ids are set in the XML directly: python-docx's style setters look the style up on every call, which made
# the SDS (tens of thousands of table cells) take minutes
def table_style(paragraph) -> None:
    paragraph._p.get_or_add_pPr().get_or_add_pStyle().val = 'FiskoTable'


def code_style(run) -> None:
    run._r.get_or_add_rPr().get_or_add_rStyle().val = 'FiskoCode'


def grid_style(table) -> None:
    table._tbl.tblPr.style = 'TableGrid'


def add_field(paragraph, instr: str, placeholder: str = '') -> None:
    """A Word field (TOC, PAGE ...). Word fills it when the fields are updated."""
    def fld(kind: str):
        r = OxmlElement('w:r')
        c = OxmlElement('w:fldChar')
        c.set(qn('w:fldCharType'), kind)
        r.append(c)
        return r
    paragraph._p.append(fld('begin'))
    r = OxmlElement('w:r')
    t = OxmlElement('w:instrText')
    t.set(qn('xml:space'), 'preserve')
    t.text = f' {instr} '
    r.append(t)
    paragraph._p.append(r)
    paragraph._p.append(fld('separate'))
    if placeholder:
        r = OxmlElement('w:r')
        t = OxmlElement('w:t')
        t.text = placeholder
        r.append(t)
        paragraph._p.append(r)
    paragraph._p.append(fld('end'))


@functools.lru_cache(maxsize=200_000)
def _thai_words(run: str) -> str:
    try:
        from pythainlp.tokenize import word_tokenize
    except ImportError:   # without the segmenter Word still opens the file; it just wraps Thai only at spaces
        return run
    return '​'.join(word_tokenize(run, engine='newmm'))


def thai_wrap(text: str) -> str:
    """Zero-width spaces between Thai words. Word breaks Thai lines only where Thai is an installed editing
    language; elsewhere a long Thai phrase moves to the next line whole and leaves short ragged lines."""
    return re.sub(r'[฀-๿]{6,}', lambda m: _thai_words(m.group(0)), text)


INLINE_RE = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)\s]+\)|(?<![*\w])\*(?!\s).+?(?<!\s)\*(?![*\w]))')


def add_inline(paragraph, text: str, lang: str, size: float | None = None, bold: bool = False) -> None:
    """The small Markdown inline dialect: **bold**, *italic*, `code`, [text](url) (kept as text).

    Fonts and sizes come from the paragraph style; size is only for the few places that differ (cover page)."""
    text = re.sub(r'<!--.*?-->', '', text)
    for part in INLINE_RE.split(text):
        if not part:
            continue
        mono = italic = strong = False
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            part, strong = part[2:-2], True
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            part, mono = part[1:-1], True
        elif part.startswith('[') and '](' in part:
            part = part[1:part.index('](')]
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            part, italic = part[1:-1], True
        run = paragraph.add_run(part if mono else thai_wrap(part))
        if mono:
            code_style(run)
        if size and not mono:
            run_size(run, lang, size)
        if bold or strong:
            run.bold = True
        if italic:
            run.italic = True


def para(doc, text: str, lang: str, size: float | None = None, bold: bool = False, align=None, after: float | None = None,
         style: str | None = None):
    p = doc.add_paragraph(style=style)
    add_inline(p, text, lang, size, bold)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    return p


def setup(doc, lang: str, code: str, version: str) -> None:
    """Page, styles (Normal, headings, Fisko Table, Fisko Code), header and footer."""
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    sz = SIZES[lang]
    base = THAI_FONT if lang == 'th' else LATIN_FONT
    styles = doc.styles
    normal = styles['Normal']
    style_font(normal, base, sz['body'], cs_size(lang, sz['body']))
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.0
    rpr = normal.element.get_or_add_rPr()
    lang_el = OxmlElement('w:lang')
    lang_el.set(qn('w:val'), 'th-TH' if lang == 'th' else 'en-GB')
    lang_el.set(qn('w:bidi'), 'th-TH')
    rpr.append(lang_el)
    for lv in range(1, 5):
        st = styles[f'Heading {lv}']
        style_font(st, base, sz['h'][lv - 1], cs_size(lang, sz['h'][lv - 1]))
        st.font.bold = True
        st.font.color.rgb = GREEN if lv <= 2 else RGBColor(0x22, 0x22, 0x22)
        st.paragraph_format.space_before = Pt(14 if lv == 1 else 10)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.0
        st.paragraph_format.keep_with_next = True
    tbl = styles.add_style('Fisko Table', WD_STYLE_TYPE.PARAGRAPH)
    tbl.base_style = normal
    style_font(tbl, base, sz['table'], cs_size(lang, sz['table']))
    tbl.paragraph_format.space_after = Pt(0)
    tbl.paragraph_format.space_before = Pt(0)
    mono = styles.add_style('Fisko Code', WD_STYLE_TYPE.CHARACTER)
    # Consolas is much larger than TH Sarabun New at the same point size
    style_font(mono, MONO_FONT, sz['code'], cs_size(lang, sz['table']))
    # header: document code and version; footer: page number
    hp = sec.header.paragraphs[0]
    hp.style = tbl
    add_inline(hp, f'{code} · {version}', lang)
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp = sec.footer.paragraphs[0]
    fp.style = tbl
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, 'PAGE', '1')
    # ask Word to update the TOC and page fields when the file is opened
    upd = OxmlElement('w:updateFields')
    upd.set(qn('w:val'), 'true')
    doc.settings.element.append(upd)


def col_widths(head: list[str], rows: list[list[str]], lang: str) -> list[float]:
    """Column widths in cm: each column at least as wide as its longest unbreakable token (ids, codes, URLs) and
    its header words, the rest shared by typical text length. Thai has no spaces and Word breaks it between words,
    so a Thai run counts as one word of at most 7 characters."""
    char_cm = 0.17 if lang == 'th' else 0.165
    mins, weights = [], []
    for k in range(len(head)):
        vals = [r[k] for r in rows]
        toks = [len(t) for v in vals for t in re.findall(r'[^\s฀-๿]+', re.sub(r'[*`]', '', v))]
        toks += [min(len(t), 7) for v in vals for t in re.findall(r'[฀-๿]+', v)]
        toks += [len(t) for t in re.sub(r'[*`]', '', head[k]).split()]
        mins.append(min(max(max(toks, default=2) * char_cm + 0.35, 1.1), 4.5))
        lens = sorted(len(v) for v in vals)
        weights.append(max(lens[int(len(lens) * 0.8)], 3) ** 0.85)
    return spread(mins, weights)


def spread(mins: list[float], weights: list[float]) -> list[float]:
    total = TEXT_W.cm
    rest = total - sum(mins)
    if rest <= 0:
        return [m * total / sum(mins) for m in mins]
    return [m + rest * w / sum(weights) for m, w in zip(mins, weights)]


def add_table(doc, head: list[str], rows: list[list[str]], lang: str, widths: list[float] | None = None,
              align: list[str] | None = None) -> None:
    """widths: relative weights, or None to size the columns from their contents."""
    cm = col_widths(head, rows, lang) if widths is None else spread([0.0] * len(widths), widths)
    t = doc.add_table(rows=0, cols=len(head))
    grid_style(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    t._tbl.tblPr.append(layout)
    # the grid first: every row added afterwards takes its cell widths from it
    for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), cm):
        gc.set(qn('w:w'), str(int(w * 567)))
    hdr = t.add_row()
    repeat_header(hdr)
    for k, c in enumerate(hdr.cells):
        p = c.paragraphs[0]
        table_style(p)
        add_inline(p, head[k], lang, bold=True)
        shade(c, 'E8F5EC')
    for r in rows:
        for k, c in enumerate(t.add_row().cells):
            p = c.paragraphs[0]
            table_style(p)
            add_inline(p, r[k], lang)
            if align and k < len(align) and align[k] == 'r':
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_image(doc, path: pathlib.Path, caption: str, lang: str) -> None:
    from PIL import Image
    data = path.read_bytes()
    img = Image.open(io.BytesIO(data))
    if img.format not in ('PNG', 'JPEG'):
        # prototype screens are WebP, which Word cannot show: JPEG keeps the file size down for screenshots
        buf = io.BytesIO()
        img.convert('RGB').save(buf, 'JPEG', quality=85)
        data = buf.getvalue()
    w, h = img.size
    width = TEXT_W
    if w and h and width * h / w > MAX_FIG_H:
        width = int(MAX_FIG_H * w / h)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(io.BytesIO(data), width=width)
    para(doc, caption, lang, align=WD_ALIGN_PARAGRAPH.CENTER, after=8, style='Fisko Table')


def add_list(doc, block: dict, pick, lang: str, level: int = 0) -> None:
    # numbered lists are numbered in the text: Word's List Number style would continue counting across lists
    ordered = block['t'] == 'ol'
    style = 'List Bullet' + (f' {min(level + 1, 3)}' if level else '')
    for n, it in enumerate(block['items'], 1):
        if ordered:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75 * (level + 1))
            p.paragraph_format.first_line_indent = Cm(-0.75)
            add_inline(p, f'{n}.\t' + pick(it['text']), lang)
        else:
            p = doc.add_paragraph(style=style)
            add_inline(p, pick(it['text']), lang)
        p.paragraph_format.space_after = Pt(2)
        if it.get('sub'):
            add_list(doc, it['sub'], pick, lang, level + 1)


def add_box(doc, fill: str):
    """A one-cell shaded table (callouts and code blocks); returns its paragraph."""
    t = doc.add_table(rows=1, cols=1)
    grid_style(t)
    c = t.rows[0].cells[0]
    shade(c, fill)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return c.paragraphs[0]


def add_callout(doc, label: str, text: str, kind: str, lang: str) -> None:
    p = add_box(doc, CALLOUT_FILL.get(kind, 'F4F5F7'))
    add_inline(p, f'**{label}** {text}' if label else text, lang)
    p.paragraph_format.space_after = Pt(0)


def add_code(doc, text: str, lang: str) -> None:
    p = add_box(doc, 'F5F5F5')
    table_style(p)
    for k, line in enumerate(text.split('\n')):
        if k:
            p.add_run().add_break(WD_BREAK.LINE)
        code_style(p.add_run(line))


# ----------------------------------------------------------------------------------------------------------
# SRS and SDS
# ----------------------------------------------------------------------------------------------------------

def cover(doc, lang: str, meta: dict) -> None:
    def L(th: str, en: str) -> str:
        return th if lang == 'th' else en
    big = SIZES[lang]['h'][0]
    para(doc, 'Fisko', lang, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, meta['title'][lang], lang, big * 1.4, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, meta['subtitle'][lang], lang, align=WD_ALIGN_PARAGRAPH.CENTER, after=30)
    add_table(doc, [L('รายการ', 'Item'), L('รายละเอียด', 'Detail')], [
        [L('รหัสเอกสาร', 'Document code'), meta['code']],
        [L('ฉบับ', 'Version'), meta['version']],
        [L('วันที่', 'Date'), meta['date']],
        [L('สถานะ', 'Status'), meta['status'][lang]],
    ], lang, widths=[1, 2.5])
    doc.add_page_break()
    para(doc, L('สารบัญ', 'Contents'), lang, SIZES[lang]['h'][0], bold=True)
    add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u',
              L('คลิกขวาแล้วเลือก "อัปเดตเขตข้อมูล" เพื่อสร้างสารบัญ', 'Right-click and choose "Update Field" to build the table of contents'))
    doc.add_page_break()


def coverage_rows(kind: str, reg: bs.Registry, law: list[dict], lang: str) -> list[list[str]]:
    mine = kind.upper()
    cov = bs.covered_refs(reg, law, kind)
    out = []
    owed = {'SRS': ('ตอบใน SRS', 'Answered in the SRS'), 'SDS': ('ตอบใน SDS', 'Answered in the SDS'),
            'PLAN': ('ตอบในแผนงาน', 'Answered in a plan'), 'COM': ('เงื่อนไขเชิงพาณิชย์', 'Commercial term'),
            'INFO': ('ข้อมูลประกอบ', 'Background')}
    li = 0 if lang == 'th' else 1
    for r in sorted(law, key=lambda r: (r['doc'], bs.ref_key(r['ref']))):
        disp = r['disposition']
        if r['ref'] in cov:
            ids = cov[r['ref']]
            state = ('ครอบคลุมแล้ว', 'Covered')[li] + ': ' + ', '.join(ids[:6]) + (f' (+{len(ids) - 6})' if len(ids) > 6 else '')
        elif disp == mine:
            state = ('ยังไม่จัดทำ', 'Not yet specified')[li]
        else:
            state = owed.get(disp, (disp, disp))[li]
        out.append([r['ref'], r['scope'], disp, r[lang], state])
    return out


def status_rows(reg: bs.Registry, prefix: str, lang: str) -> list[list[str]]:
    c = bs.status_counts(reg, prefix)
    rows, tot = [], {s: 0 for s in bs.STATUSES}
    for typ, th, en in bs.STATUS_TYPES:
        n = sum(c[typ].values())
        if not n:
            continue
        for s in bs.STATUSES:
            tot[s] += c[typ][s]
        rows.append([f'{th if lang == "th" else en} `{prefix}-{typ}`'] + [str(c[typ][s]) for s in bs.STATUSES] + [str(n)])
    rows.append(['**' + ('รวม' if lang == 'th' else 'Total') + '**'] + [f'**{tot[s]}**' for s in bs.STATUSES] + [f'**{sum(tot.values())}**'])
    return rows


def matrix_rows(reg: bs.Registry, lang: str) -> list[list[str]]:
    rows, grand = [], {s: 0 for s in bs.STATUSES}
    xc = sum(1 for k in reg.reqs if k.startswith('XC-') and bs.REQ_ID.match(k))
    rows.append(['ข้อกำหนดร่วมที่สร้างจากทะเบียนข้อกฎหมาย (XC)' if lang == 'th' else 'Cross-cutting (XC) generated from the clause register', str(xc),
                 'ไม่ประเมินสถานะ' if lang == 'th' else 'not assessed', '', ''])
    for n in list(range(1, 14)) + ['MOB']:
        prefix = 'MOB' if n == 'MOB' else f'S{n:02d}'
        c = bs.status_counts(reg, prefix)
        tot = {s: sum(c[t][s] for t in c) for s in bs.STATUSES}
        for s in bs.STATUSES:
            grand[s] += tot[s]
        label = ('แอปมือถือ (MOB)' if lang == 'th' else 'Mobile apps (MOB)') if n == 'MOB' else (f'ระบบที่ {n}' if lang == 'th' else f'System {n}')
        rows.append([label, str(sum(tot.values()))] + [str(tot[s]) for s in bs.STATUSES])
    rows.append(['**' + ('รวมทุกระบบ (ไม่รวม XC)' if lang == 'th' else 'All systems (without XC)') + '**', f'**{sum(grand.values())}**']
                + [f'**{grand[s]}**' for s in bs.STATUSES])
    return rows


def spec_docx(doc_model: dict, reg: bs.Registry, law: list[dict], lang: str) -> pathlib.Path:
    """One language of the SRS or SDS, from the blocks build_spec.py has already prepared."""
    kind, meta = doc_model['kind'], doc_model['meta']
    bs.CUR['doc'] = kind
    doc = Document()
    setup(doc, lang, meta['code'], f"v{meta['version']} · {meta['date']}")
    cover(doc, lang, meta)
    pick = (lambda d: d[lang])
    fig_word = 'รูปที่' if lang == 'th' else 'Figure'
    for b in doc_model['blocks']:
        t = b['t']
        if t == 'h':
            h = doc.add_heading(level=min(b['level'], 4))
            add_inline(h, f"{b['num']} {b['text'][lang]}", lang)
            if b['level'] == 1 and b is not doc_model['blocks'][0]:
                h.paragraph_format.page_break_before = True
        elif t == 'p':
            para(doc, b['text'][lang], lang)
        elif t in ('ul', 'ol'):
            add_list(doc, b, pick, lang)
        elif t == 'table':
            head = [x[lang] for x in b['head']]
            rows = [[c[lang] for c in r] for r in b['rows']]
            add_table(doc, head, rows, lang, None, b.get('align'))
        elif t == 'code':
            add_code(doc, b['text'][lang], lang)
        elif t == 'quote':
            lab = bs.CALLOUT_LABEL[b['kind']][0 if lang == 'th' else 1]
            add_callout(doc, lab, b['text'][lang], b['kind'], lang)
        elif t == 'hr':
            pass
        elif t == 'fig':
            path = (bs.SCREENS / lang / f"{b['id']}.webp") if b.get('kind') == 'screen' else (DIAG_PNG / lang / f"{b['id']}.png")
            caption = f"**{fig_word} {b['n']}** {b['caption'][lang]}"
            if path.exists():
                add_image(doc, path, caption, lang)
            else:
                bs.warn(f'docx: {path.relative_to(ROOT)} is missing; the figure is left as its caption')
                para(doc, caption + f' ({path.name})', lang, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif t == 'dir':
            name = b['name']
            if name == 'coverage':
                head = (['ข้อกฎหมาย', 'ระบบ', 'ตอบใน', 'ข้อกำหนด', 'สถานะ'] if lang == 'th'
                        else ['Clause', 'Scope', 'Owed to', 'Requirement', 'Status'])
                add_table(doc, head, coverage_rows(kind, reg, law, lang), lang, [2, 1, 1, 7, 3])
            elif name == 'figures':
                for fid, f in sorted(((k, v) for k, v in reg.figs.items() if v['doc'] == kind), key=lambda kv: kv[1]['n']):
                    para(doc, f"**{fig_word} {f['n']}** {f['caption'][lang]}", lang, after=1)
            elif name == 'status-summary':
                head = (['ชนิด', 'มีในต้นแบบ (Prototyped)', 'บางส่วน (Partial)', 'วางแผน (Planned)', 'รวม'] if lang == 'th'
                        else ['Type', 'Prototyped', 'Partial', 'Planned', 'Total'])
                add_table(doc, head, status_rows(reg, b['args'].get('prefix', ''), lang), lang, [4, 1.5, 1.5, 1.5, 1])
            elif name == 'status-matrix':
                head = (['ชุดข้อกำหนด', 'จำนวน', 'มีในต้นแบบ', 'บางส่วน', 'วางแผน'] if lang == 'th'
                        else ['Set', 'Rows', 'Prototyped', 'Partial', 'Planned'])
                add_table(doc, head, matrix_rows(reg, lang), lang, [4, 1, 1.5, 1.5, 1.5])
    out = bs.DIST / f"Fisko-{kind.upper()}-{lang}.docx"
    doc.save(out)
    return out


def build_spec_docx() -> list[pathlib.Path]:
    """Standalone: parse docs/spec/src again (build_spec.py passes its own parsed documents instead)."""
    law = bs.load_law()
    lock = bs.load_id_lock()
    docs = {k: bs.load_doc(k) for k in ('srs', 'sds')}
    reg = bs.Registry()
    reg.law = {r['ref']: r for r in law}
    for k in ('srs', 'sds'):
        bs.prepare(docs[k], law, lock, reg)
    bs.DIST.mkdir(parents=True, exist_ok=True)
    return [spec_docx(docs[k], reg, law, lg) for k in ('srs', 'sds') for lg in bs.LANGS]


# ----------------------------------------------------------------------------------------------------------
# ----------------------------------------------------------------------------------------------------------


if __name__ == '__main__':
    for p in build_spec_docx():
        print(p.relative_to(ROOT))
